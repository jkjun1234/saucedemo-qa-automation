"""역할: TC 엑셀의 실행 단계를 읽고, 선택한 파일에 결과·비고를 기록합니다.

TC_웹주문에서 TC 번호·시나리오를, 자동화_실행단계에서 순서·동작·대상·값을
읽습니다. read_cases()는 읽기 전용이며 과거 결과를 반환하지 않습니다.
write_results()만 결과·비고를 갱신합니다. 브라우저 조작은 step_runner 담당입니다.
"""

import os
import re
import tempfile
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook


SUPPORTED_ACTIONS = {"접속", "입력", "클릭", "URL확인", "표시확인", "최소개수확인"}
CHECK_ACTIONS = {"URL확인", "표시확인", "최소개수확인"}
TC_HEADERS = {"TC-ID", "테스트 시나리오", "결과", "비고"}
STEP_HEADERS = {"TC-ID", "순서", "동작", "대상", "입력값", "기대값", "설명"}


def _text(value):
    """빈 셀만 빈 문자열로 바꾸며 입력값의 공백은 보존합니다."""
    return "" if value is None else str(value)


def _rows(workbook, sheet_name, required_headers):
    """헤더 위치가 달라도 열 이름으로 데이터를 찾습니다."""
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"필요한 시트가 없습니다: {sheet_name}")
    rows = enumerate(workbook[sheet_name].iter_rows(values_only=True), start=1)
    for row_number, values in rows:
        headers = [_text(value).strip() for value in values]
        if "TC-ID" in headers:
            break
    else:
        raise ValueError(f"{sheet_name}: TC-ID 헤더 행을 찾지 못했습니다.")
    named_headers = [name for name in headers if name]
    if len(named_headers) != len(set(named_headers)):
        raise ValueError(f"{sheet_name} {row_number}행: 중복 헤더가 있습니다.")
    missing = required_headers - set(headers)
    if missing:
        raise ValueError(f"{sheet_name} {row_number}행: 필요한 열 누락: {', '.join(sorted(missing))}")
    for row_number, values in rows:
        if any(value is not None and value != "" for value in values):
            yield row_number, {name: values[index] for index, name in enumerate(headers) if name}


def _relative_path(value, location):
    """실행기의 대상 주소를 우회하는 절대 주소를 받지 않습니다."""
    parsed = urlsplit(value)
    if parsed.scheme or parsed.netloc:
        raise ValueError(f"{location}: 대상 주소를 기준으로 한 상대 경로를 사용하세요.")


def read_cases(path: Path) -> list[dict]:
    """준비된 TC만 번호순으로 반환하며 선택한 엑셀은 읽기 전용으로 엽니다.

    시나리오 표에만 있고 실행 단계가 없는 TC는 아직 실행 대상이 아닙니다.
    단계가 있지만 검증이 없거나 잘못된 값이 있으면 시트·행과 함께 중단합니다.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        cases = {}
        for row_number, row in _rows(workbook, "TC_웹주문", TC_HEADERS):
            tc_id = _text(row["TC-ID"]).strip()
            location = f"TC_웹주문 {row_number}행"
            if not re.fullmatch(r"TC\d+", tc_id):
                raise ValueError(f"{location}: TC-ID는 TC01 같은 형식이어야 합니다.")
            if tc_id in cases:
                raise ValueError(f"{location}: 중복 TC-ID: {tc_id}")
            cases[tc_id] = {"tc_id": tc_id, "scenario": _text(row["테스트 시나리오"]), "steps": []}

        for row_number, row in _rows(workbook, "자동화_실행단계", STEP_HEADERS):
            location = f"자동화_실행단계 {row_number}행"
            tc_id = _text(row["TC-ID"]).strip()
            if tc_id not in cases:
                raise ValueError(f"{location}: TC 표에 없는 TC-ID: {tc_id}")
            order_text = _text(row["순서"]).strip()
            if not re.fullmatch(r"[1-9]\d*", order_text):
                raise ValueError(f"{location}: 순서는 1부터 시작하는 양의 정수여야 합니다.")
            order = int(order_text)
            action = _text(row["동작"]).strip()
            target = _text(row["대상"]).strip()
            value = _text(row["입력값"])
            expected = _text(row["기대값"])
            if action not in SUPPORTED_ACTIONS:
                raise ValueError(f"{location}: 지원하지 않는 동작: {action}")
            if action in {"입력", "클릭", "표시확인", "최소개수확인"} and not target:
                raise ValueError(f"{location}: {action}에는 대상이 필요합니다.")
            if action == "접속":
                value = value.strip()
                _relative_path(value, location)
            if action == "URL확인":
                expected = expected.strip()
                if not expected:
                    raise ValueError(f"{location}: URL확인에는 기대값이 필요합니다.")
                _relative_path(expected, location)
            if action == "최소개수확인":
                if not re.fullmatch(r"[1-9]\d*", expected.strip()):
                    raise ValueError(f"{location}: 최소 개수 기대값은 양의 정수여야 합니다.")
                expected = int(expected.strip())
            if any(step["order"] == order for step in cases[tc_id]["steps"]):
                raise ValueError(f"{location}: {tc_id}의 순서 {order} 중복")
            cases[tc_id]["steps"].append({
                "order": order, "action": action, "target": target,
                "value": value, "expected": expected,
                "description": _text(row["설명"]), "source_row": row_number,
            })

        ready = []
        for tc_id in sorted(cases, key=lambda item: int(item[2:])):
            case = cases[tc_id]
            steps = sorted(case["steps"], key=lambda step: step["order"])
            if not steps:
                continue
            location = f"자동화_실행단계 {steps[0]['source_row']}행 ({tc_id})"
            if [step["order"] for step in steps] != list(range(1, len(steps) + 1)):
                raise ValueError(f"{location}: 순서는 1부터 빠짐없이 연속이어야 합니다.")
            if not any(step["action"] in CHECK_ACTIONS for step in steps):
                raise ValueError(f"{location}: 통과를 판정할 검증 단계가 없습니다.")
            case["steps"] = steps
            ready.append(case)
        if not ready:
            raise ValueError("자동화_실행단계: 준비된 실행 단계가 없습니다.")
        return ready
    finally:
        workbook.close()


def write_results(source: Path, results: list[dict]) -> Path:
    """TC-ID·헤더로 결과/비고만 갱신하고 선택한 파일 경로를 반환합니다.

    저장 중 오류가 나면 기존 파일을 유지합니다. 같은 폴더의 임시 저장 파일은
    교체 직후 제거되며 결과 사본이나 결과 시트를 남기지 않습니다.
    """
    source = Path(source)
    workbook = load_workbook(source)
    temporary = None
    try:
        rows = list(_rows(workbook, "TC_웹주문", TC_HEADERS))
        locations = {}
        for number, row in rows:
            tc_id = _text(row["TC-ID"]).strip()
            if tc_id in locations:
                raise ValueError(f"TC_웹주문 {number}행: 중복 TC-ID: {tc_id}")
            locations[tc_id] = number
        seen = set()
        for result in results:
            tc_id = result["tc_id"]
            if tc_id not in locations:
                raise ValueError(f"TC 표에 없는 TC-ID: {tc_id}")
            if tc_id in seen:
                raise ValueError(f"중복 결과 TC-ID: {tc_id}")
            if result["status"] not in {"Pass", "Fail", "N/A"}:
                raise ValueError(f"{tc_id}: 잘못된 결과값: {result['status']}")
            seen.add(tc_id)

        sheet = workbook["TC_웹주문"]
        for row in sheet:
            headers = {_text(cell.value).strip(): cell.column for cell in row if cell.value is not None}
            if "TC-ID" in headers:
                break
        for result in results:
            number = locations[result["tc_id"]]
            sheet.cell(number, headers["결과"], result["status"])
            remark = None
            if result["status"] != "Pass":
                parts = [f"사유: {result.get('reason') or '사유 미기록'}"]
                if result.get("failed_step") is not None:
                    parts.append(f"실패 단계: {result['failed_step']}")
                for key, label in [("expected", "기대값"), ("actual", "관찰/오류")]:
                    if result.get(key):
                        parts.append(f"{label}: {result[key]}")
                parts.extend(f"증거({kind}): {path}" for kind, path in result.get("evidence", {}).items() if path)
                remark = "\n".join(parts)
            sheet.cell(number, headers["비고"]).value = remark

        descriptor, name = tempfile.mkstemp(prefix=".tc-save-", suffix=".xlsx", dir=source.parent)
        os.close(descriptor)
        temporary = Path(name)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, source)
        return source
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)

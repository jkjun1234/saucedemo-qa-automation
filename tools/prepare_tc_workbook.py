"""역할: 기존 TC 초안을 보존하고 최초 실행용 v2 양식을 한 번 준비합니다.

TC01~05 정렬·실제 결과 열 제거·결과 선택 서식 이동·TC01 실행 단계를
적용합니다. 테스트나 결과 저장 프로그램이 아니며 이미 존재하는 실행용
파일은 덮어쓰지 않습니다. 원본 과거 결과는 기존 초안에 보존됩니다.
"""

from copy import copy, deepcopy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SauceDemo_TC_설계서_초안.xlsx"
DESTINATION = ROOT / "SauceDemo_TC_실행용_v2.xlsx"

STEPS = [
    ["TC01", 1, "접속", "", "", "", "프로그램에서 입력한 대상 주소의 로그인 화면으로 접속"],
    ["TC01", 2, "입력", "#user-name", "standard_user", "", "공개 데모 계정 입력"],
    ["TC01", 3, "입력", "#password", "secret_sauce", "", "공개 데모 비밀번호 입력"],
    ["TC01", 4, "클릭", "#login-button", "", "", "Login 버튼 클릭"],
    ["TC01", 5, "URL확인", "", "", "/inventory.html", "로그인 후 실제 URL과 비교"],
    ["TC01", 6, "표시확인", ".inventory_item:first-child", "", "", "첫 상품 카드가 보이는지 확인"],
    ["TC01", 7, "최소개수확인", ".inventory_item", "", 1, "상품 카드가 최소 1개인지 확인"],
]


def prepare_workbook(source: Path, destination: Path) -> Path:
    """결과 사본을 매번 만드는 기능이 아닌, 최초 실행 양식 준비 작업입니다."""
    source, destination = Path(source), Path(destination)
    if destination.exists():
        raise FileExistsError(f"기존 실행용 파일을 보존합니다: {destination}")
    workbook = load_workbook(source)
    try:
        sheet = workbook["TC_웹주문"]
        headers = [cell.value for cell in sheet[2]]
        removed_column = headers.index("실제 결과") + 1
        kept_columns = [index for index, name in enumerate(headers, 1)
                        if name and name != "실제 결과"]
        records = sorted(
            [(row[0].value, [copy(row[index - 1]) for index in kept_columns],
              sheet.row_dimensions[row[0].row].height)
             for row in sheet.iter_rows(min_row=3) if row[0].value],
            key=lambda record: int(record[0][2:]),
        )
        dimensions = [copy(sheet.column_dimensions[get_column_letter(index)])
                      for index in kept_columns]
        rules = [deepcopy(rule) for key in sheet.conditional_formatting
                 for rule in sheet.conditional_formatting[key]]
        sheet.delete_cols(removed_column)
        for column, dimension in enumerate(dimensions, 1):
            sheet.column_dimensions[get_column_letter(column)].width = dimension.width
        for row_number, (_, cells, height) in enumerate(records, 3):
            sheet.row_dimensions[row_number].height = height
            for column, source_cell in enumerate(cells, 1):
                cell = sheet.cell(row_number, column, source_cell.value)
                cell.font = copy(source_cell.font)
                cell.fill = copy(source_cell.fill)
                cell.border = copy(source_cell.border)
                cell.alignment = copy(source_cell.alignment)
                cell.number_format = source_cell.number_format
            # 새 실행용 양식에 과거 Pass나 실패 메모를 새 결과처럼 싣지 않습니다.
            for column in (8, 9):
                sheet.cell(row_number, column).value = None
                sheet.cell(row_number, column).fill = PatternFill()
                sheet.cell(row_number, column).font = copy(sheet.cell(row_number, 2).font)
        if sheet.max_column > 9:
            sheet.delete_cols(10, sheet.max_column - 9)
        sheet.merge_cells("A1:I1")
        sheet["A1"] = "실행용 양식 · TC01 단계 준비 / TC02~05 단계 미준비 · 실행 전 결과는 빈칸"
        sheet["A1"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = "D3"
        result_range = f"H3:H{len(records) + 2}"
        sheet.data_validations = DataValidationList()
        validation = DataValidation(type="list", formula1='"Pass,Fail,N/A"', allow_blank=True)
        validation.error = "Pass, Fail, N/A 중 하나를 선택하세요."
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(result_range)
        sheet.conditional_formatting = ConditionalFormattingList()
        for rule in rules:
            sheet.conditional_formatting.add(result_range, rule)

        step_sheet = workbook["자동화_실행단계"]
        step_sheet.delete_rows(1, step_sheet.max_row)
        step_sheet.append(["TC-ID", "순서", "동작", "대상", "입력값", "기대값", "설명"])
        for step in STEPS:
            step_sheet.append(step)
        for column, width in enumerate([12, 8, 19, 32, 24, 26, 57], 1):
            step_sheet.column_dimensions[get_column_letter(column)].width = width
            cell = step_sheet.cell(1, column)
            cell.font = copy(sheet["A2"].font)
            cell.fill = copy(sheet["A2"].fill)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in step_sheet.iter_rows(min_row=2):
            step_sheet.row_dimensions[row[0].row].height = 42
            for cell in row:
                cell.font = copy(sheet["B3"].font)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        step_sheet.freeze_panes = "D2"
        step_sheet.auto_filter.ref = "A1:G8"
        workbook.active = workbook.sheetnames.index("TC_웹주문")
        workbook.save(destination)
        return destination
    finally:
        workbook.close()


if __name__ == "__main__":
    print(prepare_workbook(SOURCE, DESTINATION))

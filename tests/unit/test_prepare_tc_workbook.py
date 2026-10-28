"""역할: 실행용 양식 준비 시 TC 내용·서식·원본이 보존되는지 검사합니다.

새 실행용 파일은 pytest 임시 폴더에만 만들며 실제 초안은 변경하지 않습니다.
브라우저 테스트나 결과 기록 기능의 검사가 아닙니다.
"""

from hashlib import sha256
from pathlib import Path
from copy import copy

import pytest
from openpyxl import load_workbook


SOURCE = Path(__file__).resolve().parents[2] / "SauceDemo_TC_설계서_초안.xlsx"


def test_prepares_numbered_cases_without_changing_source(tmp_path):
    from tools.prepare_tc_workbook import prepare_workbook

    before = sha256(SOURCE.read_bytes()).hexdigest()
    destination = tmp_path / "execution.xlsx"
    prepare_workbook(SOURCE, destination)
    assert sha256(SOURCE.read_bytes()).hexdigest() == before
    original = load_workbook(SOURCE)
    prepared = load_workbook(destination)
    try:
        sheet = prepared["TC_웹주문"]
        assert [sheet.cell(row, 1).value for row in range(3, 8)] == [
            "TC01", "TC02", "TC03", "TC04", "TC05",
        ]
        assert [sheet.cell(2, col).value for col in range(1, 10)] == [
            "TC-ID", "기능", "우선순위", "테스트 시나리오", "전제 조건",
            "테스트 절차", "예상 결과", "결과", "비고",
        ]
        source_cases = {row[0]: row[:7] for row in original["TC_웹주문"].iter_rows(
            min_row=3, max_row=7, values_only=True)}
        assert {row[0]: row[:7] for row in sheet.iter_rows(
            min_row=3, max_row=7, values_only=True)} == source_cases
        assert all(sheet.cell(row, col).value is None
                   for row in range(3, 8) for col in (8, 9))
        assert list(prepared["자동화_데이터"].values) == list(original["자동화_데이터"].values)
        assert copy(sheet["A2"].font) == copy(original["TC_웹주문"]["A2"].font)
    finally:
        original.close()
        prepared.close()


def test_result_controls_follow_the_result_column(tmp_path):
    from excel_io import read_cases
    from tools.prepare_tc_workbook import prepare_workbook

    destination = tmp_path / "execution.xlsx"
    prepare_workbook(SOURCE, destination)
    workbook = load_workbook(destination)
    try:
        sheet = workbook["TC_웹주문"]
        validations = sheet.data_validations.dataValidation
        assert len(validations) == 1
        assert str(validations[0].sqref) == "H3:H7"
        assert validations[0].formula1 == '"Pass,Fail,N/A"'
        assert [str(key.sqref) for key in sheet.conditional_formatting] == ["H3:H7"]
        rules = sheet.conditional_formatting["H3:H7"]
        assert len(rules) == 3
    finally:
        workbook.close()
    cases = read_cases(destination)
    assert [case["tc_id"] for case in cases] == ["TC01"]
    assert [step["order"] for step in cases[0]["steps"]] == list(range(1, 8))


def test_preparation_never_overwrites_an_existing_execution_file(tmp_path):
    from tools.prepare_tc_workbook import prepare_workbook

    destination = tmp_path / "execution.xlsx"
    destination.write_bytes(b"existing results")
    with pytest.raises(FileExistsError):
        prepare_workbook(SOURCE, destination)
    assert destination.read_bytes() == b"existing results"

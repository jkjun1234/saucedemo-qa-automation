"""역할: 임시 TC 파일에 결과·비고만 기록되고 나머지가 보존되는지 검사합니다.

개발용 검사이며 실제 사용자 TC나 SauceDemo를 실행·수정하지 않습니다.
"""

from copy import copy
from hashlib import sha256

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.xml.functions import tostring


@pytest.fixture
def source(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TC_웹주문"
    sheet.append(["유지할 안내"])
    sheet.append(["비고", "예상 결과", "TC-ID", "결과", "테스트 시나리오", "테스트 절차"])
    sheet.append(["옛 실패", "상품 표시", "TC02", "Fail", "두 번째", "기존 절차2"])
    sheet.append(["옛 실패", "오류 표시", "TC01", "Fail", "첫 번째", "기존 절차1"])
    sheet.append(["수동 기록", "기존 기대", "TC03", "Pass", "미대상", "=1+2"])
    sheet["D4"].font = Font(color="FF008000", bold=True)
    sheet.merge_cells("A1:F1")
    validation = DataValidation(type="list", formula1='"Pass,Fail,N/A"')
    sheet.add_data_validation(validation)
    validation.add("D3:D5")
    sheet.conditional_formatting.add("D3:D5", CellIsRule(operator="equal", formula=['"Pass"'], font=Font(color="FF008000")))
    steps = workbook.create_sheet("자동화_실행단계")
    steps.append(["기존 실행 단계"])
    path = tmp_path / "selected.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def result(tc_id="TC01", status="Pass"):
    return {"tc_id": tc_id, "status": status, "reason": "필수 요소 없음",
            "failed_step": 3, "actual": "요소 없음", "expected": "요소 표시",
            "evidence": {"log": "test-results/run/events.jsonl", "trace": "test-results/run/trace.zip"}}


def test_updates_same_file_by_headers_and_id_preserving_other_cells(source):
    from excel_io import write_results

    before = load_workbook(source)
    assert write_results(source, [result()]) == source
    after = load_workbook(source)
    assert after["TC_웹주문"]["D4"].value == "Pass"
    assert after["TC_웹주문"]["A4"].value is None
    assert after.sheetnames == before.sheetnames
    for name in before.sheetnames:
        for row in before[name]:
            for cell in row:
                new = after[name][cell.coordinate]
                if (name, cell.coordinate) not in {("TC_웹주문", "A4"), ("TC_웹주문", "D4")}:
                    assert new.value == cell.value
                assert copy(new.font) == copy(cell.font)
        assert str(after[name].merged_cells) == str(before[name].merged_cells)
        assert str(after[name].data_validations) == str(before[name].data_validations)
        def rules(sheet):
            return [(str(cf.sqref), [tostring(rule.to_tree()) for rule in sheet.conditional_formatting[cf]])
                    for cf in sheet.conditional_formatting]
        assert rules(after[name]) == rules(before[name])
        assert after[name].max_column == before[name].max_column
    before.close()
    after.close()
    assert sorted(p.name for p in source.parent.iterdir()) == [source.name]


@pytest.mark.parametrize("status", ["Fail", "N/A"])
def test_failure_and_na_include_reason_step_and_evidence(source, status):
    from excel_io import write_results

    write_results(source, [result(status=status)])
    workbook = load_workbook(source)
    sheet = workbook["TC_웹주문"]
    assert sheet["D4"].value == status
    for text in ["필수 요소 없음", "3", "요소 표시", "요소 없음", "events.jsonl", "trace.zip"]:
        assert text in sheet["A4"].value
    workbook.close()


@pytest.mark.parametrize("change, message", [
    ("missing", "없는 TC-ID"), ("duplicate", "중복 TC-ID"),
    ("duplicate_result", "중복 결과"), ("status", "결과값"),
    ("header", "필요한 열"),
])
def test_invalid_export_leaves_source_bytes_unchanged(source, change, message):
    from excel_io import write_results

    results = [result()]
    if change in {"duplicate", "header"}:
        workbook = load_workbook(source)
        workbook["TC_웹주문"]["C3" if change == "duplicate" else "D2"] = "TC01" if change == "duplicate" else "잘못된 열"
        workbook.save(source)
        workbook.close()
    if change == "missing":
        results.append(result("TC99"))
    if change == "duplicate_result":
        results.append(result())
    if change == "status":
        results[0]["status"] = "성공"
    digest = sha256(source.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match=message):
        write_results(source, results)
    assert sha256(source.read_bytes()).hexdigest() == digest


def test_locked_destination_is_reported_without_damaging_source(source, monkeypatch):
    import excel_io

    digest = sha256(source.read_bytes()).hexdigest()
    def locked(*args):
        raise PermissionError("파일 사용 중")
    monkeypatch.setattr(excel_io.os, "replace", locked)
    with pytest.raises(PermissionError, match="파일 사용 중"):
        excel_io.write_results(source, [result()])
    assert sha256(source.read_bytes()).hexdigest() == digest
    assert sorted(p.name for p in source.parent.iterdir()) == [source.name]

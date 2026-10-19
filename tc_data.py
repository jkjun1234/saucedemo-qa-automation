"""역할: TC 엑셀에서 테스트 실행에 필요한 데이터를 읽어 주는 모듈입니다.

자동화_데이터 시트의 A열 TC-ID로 행을 찾고 B·C·D열의 아이디·비밀번호·
예상 URL을 반환합니다. 브라우저 조작이나 결과 저장은 하지 않습니다.
tests/test_tc01_login.py가 이 함수를 사용합니다. 작성: AI, 2026-10-19.
"""

from pathlib import Path

from openpyxl import load_workbook


TC_WORKBOOK = Path(__file__).resolve().parent / "SauceDemo_TC_설계서_초안.xlsx"


def load_login_data(tc_id, workbook_path=TC_WORKBOOK):
    """자동화_데이터 시트에서 TC 번호에 해당하는 로그인 값을 읽습니다."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["자동화_데이터"]
        for row in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
            row_tc_id, username, password, expected_url = row
            if row_tc_id == tc_id:
                return {
                    "username": username,
                    "password": password,
                    "expected_url": expected_url,
                }
        raise ValueError(f"자동화_데이터 시트에서 {tc_id}를 찾지 못했습니다.")
    finally:
        workbook.close()
